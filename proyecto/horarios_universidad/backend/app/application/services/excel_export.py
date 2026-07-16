"""Builds a multi-sheet Excel workbook for a schedule run.

Produces a human-readable export mirroring the on-screen schedule grid,
plus flat sheets for subjects, teachers, groups, assignments, and solver
statistics — everything a director needs to review a run outside the app.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING, Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from app.domain.entities.academic import Subject
    from app.domain.entities.infrastructure import Classroom
    from app.domain.entities.personnel import Instructor
    from app.domain.entities.schedule import Assignment, ScheduleRun, SessionBlock
    from app.domain.entities.scheduling import Group

__all__ = ["build_schedule_workbook"]

_DAY_NAMES = ["", "Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
_START_HOUR = 7
_END_HOUR = 21

_HEADER_FILL = PatternFill("solid", fgColor="A6061B")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(bold=True, size=14, color="A6061B")
_THIN = Side(style="thin", color="D1D5DB")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CLASS_FILL = PatternFill("solid", fgColor="FCEBED")
_WRAP = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _period_to_hour(period: int) -> int:
    return _START_HOUR + period - 1


def _assign_lanes(sessions: list[dict]) -> list[dict]:
    """Assign a non-overlapping lane index to each session, sorted by start."""
    ordered = sorted(sessions, key=lambda s: (s["start_period"], s["start_period"] + s["duration"]))
    lane_end: list[int] = []
    result = []
    for s in ordered:
        start, end = s["start_period"], s["start_period"] + s["duration"]
        lane = next((i for i, e in enumerate(lane_end) if e <= start), None)
        if lane is None:
            lane = len(lane_end)
            lane_end.append(end)
        else:
            lane_end[lane] = end
        result.append({**s, "lane": lane})
    return result


def _style_header_row(ws, row: int, num_cols: int) -> None:
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autosize(ws, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def build_schedule_workbook(
    run: ScheduleRun,
    assignments: list[Assignment],
    session_blocks: list[SessionBlock],
    groups: list[Group],
    subjects: list[Subject],
    teachers: list[Instructor],
    classrooms: list[Classroom],
) -> bytes:
    subject_map = {s.id: s for s in subjects}
    teacher_map = {t.id: t for t in teachers}
    classroom_map = {c.id: c for c in classrooms}
    group_map = {g.id: g for g in groups}
    assignment_map = {a.id: a for a in assignments}

    def subject_of_group(group_id):
        group = group_map.get(group_id)
        return subject_map.get(group.subject_id) if group else None

    def teacher_name(instructor_id) -> str:
        t = teacher_map.get(instructor_id)
        return f"{t.first_name} {t.last_name}" if t else "—"

    enriched = []
    for block in session_blocks:
        assignment = assignment_map.get(block.assignment_id)
        if assignment is None:
            continue
        subject = subject_of_group(assignment.group_id)
        classroom = classroom_map.get(block.classroom_id)
        group = group_map.get(assignment.group_id)
        enriched.append({
            "day_of_week": block.day_of_week,
            "start_period": block.start_period,
            "duration": block.duration,
            "subject_name": subject.name if subject else "—",
            "subject_code": subject.code if subject else "—",
            "teacher_name": teacher_name(assignment.instructor_id),
            "classroom_name": classroom.name if classroom else "—",
            "group_number": group.group_number if group else 0,
            "group_label": f"G{group.group_number}" if group else "",
        })

    wb = Workbook()

    _build_grid_sheet(wb.active, run, enriched)
    _build_assignments_sheet(wb.create_sheet("Asignaciones"), enriched)
    _build_subjects_sheet(wb.create_sheet("Materias"), subjects)
    _build_teachers_sheet(wb.create_sheet("Docentes"), teachers)
    _build_groups_sheet(wb.create_sheet("Grupos"), groups, subject_map)
    _build_stats_sheet(wb.create_sheet("Estadísticas"), run)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _build_grid_sheet(ws, run: ScheduleRun, enriched: list[dict]) -> None:
    ws.title = "Horario"
    ws["A1"] = f"Horario — {run.name or run.id}"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + (_END_HOUR - _START_HOUR))

    header_row = 3
    ws.cell(row=header_row, column=1, value="Día")
    hours = list(range(_START_HOUR, _END_HOUR))
    for i, h in enumerate(hours):
        ws.cell(row=header_row, column=2 + i, value=f"{h}:00")
    _style_header_row(ws, header_row, 1 + len(hours))

    current_row = header_row + 1
    for day in range(1, 7):
        day_sessions = [s for s in enriched if s["day_of_week"] == day]
        laned = _assign_lanes(day_sessions)
        num_lanes = max((s["lane"] for s in laned), default=-1) + 1
        num_lanes = max(num_lanes, 1)

        day_start_row = current_row
        ws.cell(row=day_start_row, column=1, value=_DAY_NAMES[day])
        ws.cell(row=day_start_row, column=1).font = Font(bold=True)
        ws.cell(row=day_start_row, column=1).alignment = Alignment(vertical="center")
        if num_lanes > 1:
            ws.merge_cells(start_row=day_start_row, start_column=1, end_row=day_start_row + num_lanes - 1, end_column=1)

        for lane in range(num_lanes):
            row = day_start_row + lane
            for col in range(2, 2 + len(hours)):
                ws.cell(row=row, column=col).border = _BORDER

        for s in laned:
            row = day_start_row + s["lane"]
            start_col = 2 + (s["start_period"] - 1)
            end_col = start_col + s["duration"] - 1
            end_col = min(end_col, 1 + len(hours))
            cell = ws.cell(row=row, column=start_col)
            cell.value = (
                f"{s['subject_name']} ({s['subject_code']})\n"
                f"{s['teacher_name']}\n"
                f"{s['classroom_name']} · {s['group_label']}"
            )
            cell.fill = _CLASS_FILL
            cell.alignment = _WRAP
            cell.border = _BORDER
            if end_col > start_col:
                ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)

        current_row = day_start_row + num_lanes

    _autosize(ws, [12] + [14] * len(hours))
    for r in range(header_row, current_row):
        ws.row_dimensions[r].height = 48


def _build_assignments_sheet(ws, enriched: list[dict]) -> None:
    headers = ["Día", "Inicio", "Duración (periodos)", "Materia", "Código", "Docente", "Salón", "Grupo"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    _style_header_row(ws, 1, len(headers))

    rows = sorted(enriched, key=lambda s: (s["day_of_week"], s["start_period"]))
    for r, s in enumerate(rows, start=2):
        values = [
            _DAY_NAMES[s["day_of_week"]],
            f"{_period_to_hour(s['start_period'])}:00",
            s["duration"],
            s["subject_name"],
            s["subject_code"],
            s["teacher_name"],
            s["classroom_name"],
            s["group_label"],
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = _BORDER

    _autosize(ws, [12, 10, 16, 28, 12, 24, 16, 10])


def _build_subjects_sheet(ws, subjects: list[Subject]) -> None:
    headers = ["Código", "Nombre", "Créditos", "Horas semanales", "Modalidad", "Estado"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    _style_header_row(ws, 1, len(headers))

    for r, s in enumerate(subjects, start=2):
        values = [s.code, s.name, s.credit_hours, s.weekly_hours, s.modality_default, "Activa" if s.is_active else "Inactiva"]
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v).border = _BORDER

    _autosize(ws, [12, 32, 10, 16, 14, 12])


def _build_teachers_sheet(ws, teachers: list[Instructor]) -> None:
    headers = ["Código", "Nombre", "Apellido", "Correo", "Teléfono", "Estado"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    _style_header_row(ws, 1, len(headers))

    for r, t in enumerate(teachers, start=2):
        values = [t.employee_code, t.first_name, t.last_name, t.email, t.phone, "Activo" if t.is_active else "Inactivo"]
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v).border = _BORDER

    _autosize(ws, [12, 18, 18, 28, 16, 12])


def _build_groups_sheet(ws, groups: list[Group], subject_map: dict) -> None:
    headers = ["Materia", "Código", "Grupo #", "Cupo esperado", "Modalidad", "Prioridad"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    _style_header_row(ws, 1, len(headers))

    for r, g in enumerate(groups, start=2):
        subject = subject_map.get(g.subject_id)
        values = [
            subject.name if subject else "—",
            subject.code if subject else "—",
            g.group_number,
            g.expected_enrollment,
            g.modality,
            g.priority,
        ]
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v).border = _BORDER

    _autosize(ws, [28, 12, 10, 14, 14, 10])


def _build_stats_sheet(ws, run: ScheduleRun) -> None:
    ws["A1"] = "Estadísticas del solver"
    ws["A1"].font = _TITLE_FONT

    rows: list[tuple[str, Any]] = [
        ("Nombre de la corrida", run.name),
        ("Estado", run.status),
        ("Iniciado", run.started_at.isoformat() if run.started_at else ""),
        ("Completado", run.completed_at.isoformat() if run.completed_at else ""),
        ("Valor objetivo", run.objective_value),
        ("Gap", run.gap),
    ]

    stats = run.statistics or {}
    for key in (
        "groups_total", "groups_scheduled", "instructors_total", "rooms_total",
        "time_slots_total", "num_variables", "num_constraints",
        "hard_constraints_total", "soft_constraints_total", "elapsed_seconds",
        "memory_usage_mb", "workload_spread", "avg_preference_score",
        "room_utilization_pct",
    ):
        if key in stats:
            rows.append((key, stats[key]))

    r = 3
    for label, value in rows:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value=value)
        r += 1

    r += 1
    timings = stats.get("timings")
    if isinstance(timings, dict):
        ws.cell(row=r, column=1, value="Tiempos por etapa (segundos)").font = Font(bold=True)
        r += 1
        for stage, seconds in timings.items():
            ws.cell(row=r, column=1, value=stage)
            ws.cell(row=r, column=2, value=seconds)
            r += 1

    r += 1
    constraints_by_type = stats.get("constraints_by_type")
    if isinstance(constraints_by_type, dict):
        ws.cell(row=r, column=1, value="Restricciones por tipo").font = Font(bold=True)
        r += 1
        for ctype, count in constraints_by_type.items():
            ws.cell(row=r, column=1, value=ctype)
            ws.cell(row=r, column=2, value=count)
            r += 1

    _autosize(ws, [28, 20])
